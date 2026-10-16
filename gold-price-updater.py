import math
from datetime import datetime
import openpyxl
import requests


def update_gold_excel():
    try:
        url = "https://api.chnwt.dev/thai-gold-api/latest"
        response = requests.get(url, timeout=10)

        if response.status_code == 200:
            data = response.json()

            if data.get("status") == "success":
                resp_data = data["response"]
                price_data = resp_data["price"]

                # แปลงค่าเป็น float โดยตัดคอมมาออก
                gold_bar_sell = float(
                    str(price_data["gold_bar"]["sell"]).replace(",", "")
                )
                gold_bar_buy = float(
                    str(price_data["gold_bar"]["buy"]).replace(",", "")
                )
                gold_ornament_sell = float(
                    str(price_data["gold"]["sell"]).replace(",", "")
                )
                gold_ornament_buy = math.ceil(gold_bar_buy * 0.95)

                # คำนวณส่วนต่าง (Spread) ซื้อ-ขาย
                spread_ornament = gold_ornament_sell - gold_ornament_buy
                spread_bar = gold_bar_sell - gold_bar_buy

                update_str = f"อัปเดตเมื่อ: {resp_data.get('update_date', '')} {resp_data.get('update_time', '')}"

                # โหลดไฟล์ Excel
                wb = openpyxl.load_workbook("controller.xlsx")
                ws = wb.active

                # ==========================================
                # 1. อัปเดตราคาปัจจุบันไว้ที่ A1:B5 (สำหรับหน้าเว็บ)
                # ==========================================
                ws["A1"] = "รูปพรรณ ขาย"
                ws["B1"] = f"{gold_ornament_sell:,.2f}"

                ws["A2"] = "รูปพรรณ รับ"
                ws["B2"] = f"{gold_ornament_buy:,.2f}"

                ws["A3"] = "ทองแท่ง ขาย"
                ws["B3"] = f"{gold_bar_sell:,.2f}"

                ws["A4"] = "ทองแท่ง รับ"
                ws["B4"] = f"{gold_bar_buy:,.2f}"

                ws["A5"] = "เวลาอัปเดต"
                ws["B5"] = update_str

                # ==========================================
                # 2. ตารางเติมข้อมูลประวัติรายวัน (คอลัมน์ H ถึง N)
                # ==========================================
                last_row_daily = 2
                while ws[f"H{last_row_daily + 1}"].value is not None:
                    last_row_daily += 1

                should_append = True
                if last_row_daily >= 3:
                    last_ornament_sell = float(ws[f"I{last_row_daily}"].value or 0)
                    last_ornament_buy = float(ws[f"J{last_row_daily}"].value or 0)
                    last_bar_sell = float(ws[f"K{last_row_daily}"].value or 0)
                    last_bar_buy = float(ws[f"L{last_row_daily}"].value or 0)

                    # ถ้าทั้ง 4 ราคายังเท่าเดิม ไม่เพิ่มแถวซ้ำ
                    if (
                        last_ornament_sell == gold_ornament_sell
                        and last_ornament_buy == gold_ornament_buy
                        and last_bar_sell == gold_bar_sell
                        and last_bar_buy == gold_bar_buy
                    ):
                        should_append = False

                current_datetime = datetime.now()
                current_date_str = current_datetime.strftime("%Y-%m-%d %H:%M:%S")

                if should_append:
                    next_row_daily = last_row_daily + 1 if last_row_daily >= 2 else 3
                    ws[f"H{next_row_daily}"] = current_date_str
                    ws[f"I{next_row_daily}"] = gold_ornament_sell
                    ws[f"J{next_row_daily}"] = gold_ornament_buy
                    ws[f"K{next_row_daily}"] = gold_bar_sell
                    ws[f"L{next_row_daily}"] = gold_bar_buy
                    ws[f"M{next_row_daily}"] = spread_ornament
                    ws[f"N{next_row_daily}"] = spread_bar
                    print("บันทึกราคาใหม่ลงตารางรายวันสำเร็จ - gold-price-updater.py:94")

                # ==========================================
                # 3. ระบบเช็กตัดรอบรายเดือน (ทุกวันที่ 1 ของเดือน)
                # ==========================================
                # สมมติเราเช็กจากวันปัจจุบัน ถ้าเป็นวันที่ 1 และยังไม่เคยสรุปของเดือนก่อนหน้า
                # เราสามารถเขียนเงื่อนไขตรวจจับและโยกย้ายข้อมูลมาหาค่าเฉลี่ยใส่คอลัมน์ O ถึง U ได้ที่นี่ครับ
                today_day = current_datetime.day
                current_month_str = current_datetime.strftime("%Y-%m")

                # ตัวอย่างโครงสร้างการเก็บบันทึกค่าเฉลี่ยรายเดือน (สามารถพัฒนาต่อยอดเงื่อนไขวันที่ 1 ได้ตามต้องการ)
                print(f"รันระบบสำเร็จ ณ วันที่ {current_date_str}")

                # บันทึกไฟล์กลับ
                wb.save("controller.xlsx")

            else:
                print("API แจ้งเตือนสถานะ: - gold-price-updater.py:113", data)
        else:
            print(f"เชื่อมต่อ API ไม่สำเร็จ รหัสสถานะ: {response.status_code}")

    except Exception as e:
        print("เกิดข้อผิดพลาด: - gold-price-updater.py:120", e)


if __name__ == "__main__":
    update_gold_excel()
